"""Bake-off for everyday assistant work: which model should MBUDDY escalate to, and what can the local 4B do?

Synthetic tasks only (no personal data), graded by code. Cloud models run one request at a time (the Free plan allows one).
  .venv/Scripts/python.exe tools/bench_cloud.py --backend cloud --models gpt-oss:120b nemotron-3-ultra
  .venv/Scripts/python.exe tools/bench_cloud.py --backend local --models qwen3:4b-instruct
  .venv/Scripts/python.exe tools/bench_cloud.py --report
Results are appended to docs/audit/bench/cloud_runs.jsonl (a finished model/task pair is not run twice).
"""
from __future__ import annotations

import argparse
import ast
import base64
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import time
from collections import defaultdict
from pathlib import Path

import requests
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
OUT = ROOT / "docs" / "audit" / "bench" / "cloud_runs.jsonl"
CLOUD = os.getenv("OLLAMA_CLOUD_BASE_URL", "https://ollama.com").rstrip("/")
KEY = os.getenv("OLLAMA_CLOUD_API_KEY", "").strip().strip('"')
LOCAL = "http://127.0.0.1:11434"
SYSTEM = "You are a precise assistant. Follow the requested output format exactly and reply with only what was asked."
# USD per million tokens (input, output) from ollama.com/pricing, to show what the same work would cost on a paid plan
PRICE = {"gpt-oss:120b": (0.15, 0.60), "gpt-oss:20b": (0.07, 0.30), "gemma4:31b": (0.14, 0.40),
         "nemotron-3-super": (0.015, 0.60), "nemotron-3-ultra": (0.10, 3.00)}


# ───────────────────────────── helpers ─────────────────────────────
def jparse(text: str):
    t = (text or "").strip()
    m = re.search(r"```(?:json)?\s*(.*?)```", t, re.S)
    if m:
        t = m.group(1).strip()
    try:
        return json.loads(t)
    except ValueError:
        pass
    for open_c, close_c in (("{", "}"), ("[", "]")):
        a, b = t.find(open_c), t.rfind(close_c)
        if 0 <= a < b:
            try:
                return json.loads(t[a:b + 1])
            except ValueError:
                continue
    return None


def frac(*conds) -> tuple[float, str]:
    got = [bool(c) for c in conds]
    return sum(got) / len(got), "".join("1" if g else "0" for g in got)


ALLOWED = {"docx", "bs4", "re", "json", "datetime", "math", "typing", "collections", "itertools"}
BANNED_NAMES = {"open", "exec", "eval", "compile", "__import__", "input", "breakpoint", "globals", "locals"}


def safe_code(code: str) -> str:
    try:
        tree = ast.parse(code)
    except SyntaxError as exc:
        return f"syntax error: {exc.msg}"
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for a in node.names:
                if a.name.split(".")[0] not in ALLOWED:
                    return f"import {a.name} not allowed"
        elif isinstance(node, ast.ImportFrom) and (node.module or "").split(".")[0] not in ALLOWED:
            return f"import from {node.module} not allowed"
        elif isinstance(node, ast.Name) and node.id in BANNED_NAMES:
            return f"{node.id} not allowed"
    return ""


def code_of(text: str) -> str:
    m = re.search(r"```(?:python)?\s*\n(.*?)```", text or "", re.S)
    return (m.group(1) if m else text or "").strip()


def run_code(code: str, tail: str = "") -> tuple[bool, str, Path]:
    why = safe_code(code)
    folder = Path(tempfile.mkdtemp(prefix="bench_"))
    if why:
        return False, why, folder
    (folder / "run.py").write_text(code + "\n" + tail, encoding="utf-8")
    try:
        r = subprocess.run([sys.executable, "run.py"], cwd=folder, capture_output=True, text=True, timeout=40,
                           creationflags=0x08000000 if os.name == "nt" else 0)
    except subprocess.TimeoutExpired:
        return False, "timeout", folder
    return r.returncode == 0, (r.stdout + r.stderr)[-400:], folder


def png_b64(draw) -> str:
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (480, 200), "white")
    d = ImageDraw.Draw(img)
    draw(d, ImageFont)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return base64.b64encode(buf.getvalue()).decode()


# ───────────────────────────── tasks ─────────────────────────────
TOOLS = """Tools (reply with ONLY one JSON object {"tool": name, "args": {...}}):
- set_timer(minutes: number)
- add_reminder(text: string, when: ISO 8601 local datetime like 2026-10-09T15:30)
- send_email(to: string, subject: string, body: string)
- search_files(query: string, folder: string)
- open_app(name: string)
- music_play(query: string)
- calendar_add(title: string, start: ISO 8601 local datetime, duration_min: number)
- word_write(title: string, text: string)
- none(reply: string)   -- when no tool is needed, answer the user in args.reply
Today is Thursday 2026-10-08.
User: """


def tool_task(request, check):
    def chk(text):
        j = jparse(text)
        if not isinstance(j, dict) or not isinstance(j.get("args", {}), dict):
            return 0.0, "not a tool JSON"
        return check(str(j.get("tool", "")).strip(), j.get("args") or {})
    return TOOLS + request, chk


def _t1(tool, a):
    return frac(tool == "add_reminder", "dentist" in str(a.get("text", "")).lower(), str(a.get("when", "")).startswith("2026-10-09T15:30"))


def _t2(tool, a):
    try:
        ok = abs(float(a.get("minutes")) - 1.5) < 0.01
    except (TypeError, ValueError):
        ok = False
    return frac(tool == "set_timer", ok)


def _t3(tool, a):
    return frac(tool == "send_email", str(a.get("to", "")).lower() == "sara@acme.com",
                "invoice" in str(a.get("subject", "")).lower(), "attach" in str(a.get("body", "")).lower())


def _t4(tool, a):
    return frac(tool == "search_files", "tax" in str(a.get("query", "")).lower(), str(a.get("folder", "")).strip("\\/ ").lower() == "documents")


def _t5(tool, a):
    try:
        dur = float(a.get("duration_min")) == 45
    except (TypeError, ValueError):
        dur = False
    return frac(tool == "calendar_add", "omar" in str(a.get("title", "")).lower(), str(a.get("start", "")).startswith("2026-10-12T10:00"), dur)


def _t6(tool, a):
    return frac(tool == "none", "paris" in str(a.get("reply", "")).lower())


def _t7(tool, a):
    q = str(a.get("query", "")).lower()
    return frac(tool == "music_play", any(n in q for n in ("fairuz", "fairouz", "fayrouz", "feyrouz")))


def _t8(tool, a):
    return frac(tool == "word_write", "thanks" in str(a.get("title", "")).lower(), len(str(a.get("text", ""))) > 20)


EMAIL1 = ("From: Priya Nair <priya.nair@northwind.example>\nDate: Tue, 6 Oct 2026 09:12\nSubject: Offsite planning\n\n"
          "Hi team, the offsite is on Friday 23 October 2026 at 9:30 AM at the Harbor Hotel, Lisbon. Please (1) confirm your attendance "
          "by Oct 12, (2) send dietary needs to Marta, (3) book your own train tickets. Thanks, Priya")


def chk_e1(text):
    j = jparse(text) or {}
    ai = j.get("action_items")
    return frac(str(j.get("sender_name", "")).strip() == "Priya Nair", str(j.get("sender_email", "")).lower() == "priya.nair@northwind.example",
                j.get("event_date") == "2026-10-23", j.get("event_time") == "09:30", "harbor hotel" in str(j.get("location", "")).lower(),
                j.get("deadline") == "2026-10-12", isinstance(ai, list) and len(ai) == 3)


THREAD = ("Message 1 (Lina): The vendor contract needs to be signed by 17 October. I will own the paperwork.\n"
          "Message 2 (Omar): Budget is capped at $4,200 for the first year, so no add-ons.\n"
          "Message 3 (Lina): Agreed. Finance wants the signed copy back before the deadline, so please route comments to me.")


def chk_e2(text):
    lines = [l.strip() for l in (text or "").strip().splitlines() if l.strip()]
    low = " ".join(lines).lower()
    return frac(len(lines) == 3 and all(l.startswith(("-", "*", "•")) for l in lines), "17" in low and "oct" in low, "4,200" in low or "4200" in low, "lina" in low)


def chk_e3(text):
    t = text or ""
    low = t.lower()
    return frac(len(t.split()) <= 80, "tuesday" in low, bool(re.search(r"\b2(:00)?\s?(pm|p\.m\.)", low)), "!" not in t,
                "friday" in low and bool(re.search(r"sorry|apolog|unfortunately|regret", low)))


E4 = ["URGENT: server down in production, customers affected", "You won a free cruise! Click here now", "Lunch menu for next week",
      "Invoice 8841 overdue, final notice before suspension", "50% off sneakers just for you", "Minutes from yesterday's standup"]
E4_TRUTH = ["urgent", "spam", "normal", "urgent", "spam", "normal"]


def chk_e4(text):
    j = jparse(text)
    if isinstance(j, list):
        j = {str(i + 1): v for i, v in enumerate(j)}
    if not isinstance(j, dict):
        return 0.0, "no JSON"
    got = [str(j.get(str(i + 1), "")).lower().strip() for i in range(6)]
    s = sum(g == t for g, t in zip(got, E4_TRUTH)) / 6
    return s, "".join("1" if g == t else "0" for g, t in zip(got, E4_TRUTH))


def chk_e5(text):
    t = re.sub(r"[   ]", " ", text or "")      # some models use a no-break space
    low = t.lower()
    return frac(not re.search(r"\b\w+'(s|t|re|ve|ll|d|m)\b", t), "8841" in t, bool(re.search(r"oct(ober)?\.? 20|20(th)? (of )?oct", low)),
                not re.search(r"\b(hey|pls|thx|heads up)\b", low))


def chk_o1(text):
    ok, out, folder = run_code(code_of(text))
    f = folder / "report.docx"
    if not ok or not f.exists():
        return 0.0, "run failed: " + out[-90:].replace("\n", " ")
    try:
        import docx
        d = docx.Document(str(f))
    except Exception as exc:  # noqa: BLE001
        return 0.0, f"bad docx: {exc}"
    heads = [p for p in d.paragraphs if p.style.name.lower().startswith("heading 1") and p.text.strip() == "Q3 Summary"]
    para = any("Revenue grew 12% quarter over quarter." in p.text for p in d.paragraphs)
    t = d.tables[0] if d.tables else None
    cells = [[c.text.strip() for c in r.cells] for r in t.rows] if t else []
    return frac(bool(heads), para, t is not None and len(cells) == 3 and all(len(r) == 2 for r in cells),
                cells[:1] == [["Region", "Revenue"]], cells[1:] == [["North", "1.2M"], ["South", "0.9M"]])


def norm_formula(t: str) -> str:
    return re.sub(r"\s+", "", code_of(t).strip().strip("`")).upper()


def chk_o2(text):
    f = norm_formula(text)
    a = re.fullmatch(r'=SUMIF\(A2:A20,"NORTH",B2:B20\)', f)
    b = re.fullmatch(r'=SUMIFS\(B2:B20,A2:A20,"NORTH"\)', f)
    return (1.0, "ok") if (a or b) else (0.0, f[:60])


def chk_o3(text):
    f = norm_formula(text)
    return (1.0, "ok") if re.fullmatch(r"=ROUND\(AVERAGE\(C2:C50\),1\)", f) else (0.0, f[:60])


def chk_o4(text):
    j = jparse(text) or {}
    s = j.get("slides") if isinstance(j, dict) else j
    if not isinstance(s, list):
        return 0.0, "no slides"
    titles = [str(x.get("title", "")).strip() for x in s if isinstance(x, dict)]
    return frac(len(s) == 4, all(isinstance(x, dict) and isinstance(x.get("bullets"), list) and len(x["bullets"]) == 3 for x in s),
                len(set(titles)) == len(titles) and all(titles))


def chk_o5(text):
    t = text or ""
    low = t.lower()
    sec = {k: re.search(rf"{k}[^\n]*\n(.*?)(?=\n#+ |\n\*\*|\Z)", t, re.S | re.I) for k in ("attendees", "decisions", "action items")}
    a = sec["attendees"].group(1).lower() if sec["attendees"] else ""
    d = sec["decisions"].group(1).lower() if sec["decisions"] else ""
    x = sec["action items"].group(1).lower() if sec["action items"] else ""
    return frac(all(sec.values()), all(n in a for n in ("ana", "ben", "chloe")), "nov" in d and "4" in d,
                "ben" in x and "login" in x, "chloe" in x and "venue" in x, "ana" in x and "client" in x)


HTML1 = ("<table><tr><th>Product</th><th>Qty</th><th>Price</th></tr><tr><td>Pen</td><td>10</td><td>1.50</td></tr>"
         "<tr><td>Notebook</td><td>3</td><td>4.25</td></tr><tr><td>Stapler</td><td>1</td><td>12.00</td></tr></table>")


def chk_s1(text):
    j = jparse(text)
    want = [("Pen", 10, 1.5), ("Notebook", 3, 4.25), ("Stapler", 1, 12.0)]
    if not isinstance(j, list) or len(j) != 3:
        return 0.0, "not a 3-item list"
    ok = 0
    for row, (n, q, p) in zip(j, want):
        try:
            ok += isinstance(row, dict) and row.get("Product") == n and int(row.get("Qty")) == q and abs(float(row.get("Price")) - p) < 1e-6
        except (TypeError, ValueError):
            pass
    return ok / 3, f"{ok}/3 rows"


HTML2 = ('<div class="product"><h3>Desk Lamp</h3><span class="price">$19.99</span><span class="stock">In stock</span></div>'
         '<div class="product"><h3>Monitor Arm</h3><span class="price">$54.50</span><span class="stock">Out of stock</span></div>'
         '<div class="product"><h3>USB Hub</h3><span class="price">$12.00</span><span class="stock">In stock</span></div>')


def chk_s2(text):
    j = jparse(text)
    want = [("Desk Lamp", 19.99, True), ("Monitor Arm", 54.5, False), ("USB Hub", 12.0, True)]
    if not isinstance(j, list) or len(j) != 3:
        return 0.0, "not a 3-item list"
    ok = 0
    for row, (n, p, s) in zip(j, want):
        try:
            ok += isinstance(row, dict) and row.get("name") == n and abs(float(str(row.get("price")).lstrip("$")) - p) < 1e-6 and row.get("in_stock") is s
        except (TypeError, ValueError):
            pass
    return ok / 3, f"{ok}/3 rows"


HTML3 = ('<html><body><nav><a href="/home"> Home </a><a href="/about">About us</a><ul><li><a href="/pricing">Pricing</a></li></ul></nav>'
         '<main><a href="/ignore">not in nav</a></main></body></html>')


def chk_s3(text):
    tail = f'\nimport json\nprint("RESULT:" + json.dumps([list(t) for t in extract({HTML3!r})]))\n'
    ok, out, _ = run_code(code_of(text), tail)
    m = re.search(r"RESULT:(.*)", out if ok else "")
    if not ok or not m:
        return 0.0, "run failed: " + out[-90:].replace("\n", " ")
    try:
        got = json.loads(m.group(1))
    except ValueError:
        return 0.0, "bad output"
    return (1.0, "ok") if got == [["Home", "/home"], ["About us", "/about"], ["Pricing", "/pricing"]] else (0.0, str(got)[:80])


MESSY = ("Contact Ana at ana.ruiz@mail.example or call +34 912 345 678. For billing write to billing@corp.example; "
         "our office line is (415) 555-0134. Ignore the fake one: not-an-email@ and 12-34.")


def chk_s4(text):
    j = jparse(text) or {}
    em = {str(e).lower().strip() for e in j.get("emails", [])} if isinstance(j, dict) else set()
    ph = {re.sub(r"\D", "", str(p)) for p in j.get("phones", [])} if isinstance(j, dict) else set()
    return frac(em == {"ana.ruiz@mail.example", "billing@corp.example"}, ph == {"34912345678", "4155550134"} or ph == {"912345678", "4155550134"})


TASKLIST = ("Image Name        PID   Mem Usage\nchrome.exe       4120   312,456 K\nchrome.exe       4188   201,332 K\nchrome.exe       4260   150,004 K\n"
            "chrome.exe       4311    98,220 K\ncode.exe         5120   640,112 K\nexplorer.exe      880   120,500 K\npython.exe       7010    88,000 K")


def chk_w1(text):
    j = jparse(text) or {}
    try:
        mb = float(j.get("chrome_total_mb"))
    except (TypeError, ValueError):
        mb = -1
    return frac(str(j.get("top_process", "")).lower() == "code.exe", abs(mb - 744.2) < 1.5 or abs(mb - 762.0) < 1.5)


SYSINFO = ("OS Name:                   Microsoft Windows 11 Pro\nTotal Physical Memory:     16,236 MB\nAvailable Physical Memory: 5,412 MB\n"
           "System Boot Time:          10/3/2026, 8:15:00 AM")


def chk_w2(text):
    j = jparse(text) or {}
    def num(k):
        try:
            return float(j.get(k))
        except (TypeError, ValueError):
            return -1.0
    return frac("windows 11 pro" in str(j.get("os", "")).lower(), 15.7 <= num("ram_gb") <= 16.3, 5.2 <= num("free_ram_gb") <= 5.5, abs(num("uptime_days") - 5) < 0.2)


FILES = ("Name              Length      LastWriteTime\nreport_final.docx 7340032     2026-09-15\nbackup.zip        52428800    2026-08-20\n"
         "notes.txt         2048        2026-10-01\nvideo.mp4         104857600   2026-09-30\nphoto.jpg         6291456     2026-10-02\n"
         "old.iso           734003200   2026-05-05\nslides.pptx       4194304     2026-09-20")


def chk_w3(text):
    j = jparse(text)
    return (1.0, "ok") if j == ["video.mp4", "report_final.docx", "photo.jpg"] else (0.0, str(j)[:70])


UITREE = '[1] window "Untitled - Notepad"\n[2] menu "File"\n[3] menuitem "New"\n[4] menuitem "Save"\n[5] menuitem "Save As..."\n[6] menuitem "Print..."\n[7] button "Close"\n[8] textbox "Text editor"'


def chk_v1(text):
    j = jparse(text) or {}
    return (1.0, "ok") if str(j.get("id")) == "5" else (0.0, str(j)[:50])


COMPOSE = '[10] button "Attach" enabled\n[11] button "Send" disabled (needs a recipient)\n[12] textbox "To" empty\n[13] textbox "Subject" empty\n[14] textbox "Body" empty'


def chk_v2(text):
    j = jparse(text)
    steps = j.get("steps") if isinstance(j, dict) else j
    if not isinstance(steps, list) or not steps:
        return 0.0, "no steps"
    first, last = steps[0], steps[-1]
    return frac(isinstance(first, dict) and str(first.get("id")) == "12" and str(first.get("action")).lower() == "type" and "bob@x.com" in str(first.get("text", "")),
                isinstance(last, dict) and str(last.get("id")) == "11" and str(last.get("action")).lower() == "click")


OCR = "1NV0ICE  #8841\nBill To: Northwind Ltd\nSubtotal 1,100.00\nTax (17%) 187.OO\nTOTAL DUE  1,287.00\nDue date: 20 0ct 2026"


def chk_v3(text):
    j = jparse(text) or {}
    try:
        tot = abs(float(str(j.get("total")).replace(",", "")) - 1287.0) < 0.01
    except ValueError:
        tot = False
    return frac(str(j.get("invoice_number", "")).lstrip("#") == "8841", tot, j.get("due_date") == "2026-10-20")


def chk_v4(text):
    return (1.0, "ok") if re.search(r"2[, ]?450\.75", text or "") else (0.0, (text or "")[:40])


def chk_v5(text):
    return (1.0, "ok") if "green" in (text or "").lower() and "red" not in (text or "").lower() else (0.0, (text or "")[:40])


def chk_r1(text):
    j = jparse(text)
    if not isinstance(j, list) or len(j) != 4:
        return 0.0, "not 4 meetings"
    try:
        m = {x["name"].strip().upper(): (int(x["start"][:2]) * 60 + int(x["start"][3:5]), int(x["end"][:2]) * 60 + int(x["end"][3:5])) for x in j}
    except (KeyError, ValueError, TypeError, AttributeError):
        return 0.0, "bad fields"
    if set(m) != {"A", "B", "C", "D"}:
        return 0.0, "names"
    dur = {"A": 60, "B": 30, "C": 45, "D": 30}
    order = sorted(m, key=lambda k: m[k][0])
    gaps = all(m[b][0] - m[a][1] >= 15 for a, b in zip(order, order[1:]))
    return frac(all(m[k][1] - m[k][0] == dur[k] for k in m), gaps, all(540 <= m[k][0] and m[k][1] <= 780 for k in m),
                m["B"][0] < m["A"][0], order[0] != "C", order[-1] == "D")


def chk_r2(text):
    lines = [l.strip() for l in (text or "").strip().splitlines() if l.strip()]
    ok = 0
    for n in range(1, 6):
        for l in lines:
            m = re.fullmatch(rf"{n}\.\s+([a-z]+)", l)
            if m and m.group(1).startswith("b") and len(m.group(1)) == n + 2:
                ok += 1
                break
    return (ok / 5 if len(lines) == 5 else ok / 10), f"{ok}/5 lines, {len(lines)} lines"


def chk_r3(text):
    return (1.0, "ok") if "11:22" in (text or "") else (0.0, (text or "")[-40:].replace("\n", " "))


def _long_doc() -> str:
    topics = ["logistics", "procurement", "onboarding", "facilities", "compliance", "marketing", "support", "finance"]
    paras = []
    for i in range(70):
        t = topics[i % 8]
        paras.append(f"Section {i + 1} ({t}). The {t} team reviewed its weekly process and found that handoffs between departments still take "
                     f"longer than planned. Several tickets were reopened because the original notes were incomplete, and the group agreed to keep "
                     f"a shared checklist for {t} work. No budget changes were requested in this section, and the next review is scheduled as usual. "
                     f"Staff feedback was collected through the usual survey, with {40 + i % 17} responses, mostly neutral. Minor wording fixes were "
                     f"made to the {t} handbook, and a reminder was sent about keeping documents in the shared folder rather than in email.")
    paras[18] += " Note: the project budget code for the Lisbon offsite is ZX-4471."
    paras[41] += " Note: the vendor contract must be returned by March 9."
    paras[63] += " Note: the kickoff meeting will be held in room Cedar."
    return "\n\n".join(paras)


def chk_l1(text):
    j = jparse(text) or {}
    low = json.dumps(j).lower() if j else (text or "").lower()
    return frac("zx-4471" in low, "march 9" in low or "9 march" in low, "cedar" in low)


def vision(draw):
    return png_b64(draw)


def _draw_total(d, fonts):
    try:
        f = fonts.truetype("arial.ttf", 44)
    except OSError:
        f = fonts.load_default()
    d.text((30, 70), "TOTAL: 2,450.75", fill="black", font=f)


def _draw_circles(d, _fonts):
    for i, c in enumerate(("blue", "red", "green")):
        d.ellipse((40 + i * 140, 50, 140 + i * 140, 150), fill=c)


def build_tasks():
    T = [
        ("T1", "tools", *tool_task("remind me to call the dentist tomorrow at 3:30pm", _t1)),
        ("T2", "tools", *tool_task("set a timer for ninety seconds", _t2)),
        ("T3", "tools", *tool_task("email Sara at sara@acme.com that the invoice is attached, subject Invoice", _t3)),
        ("T4", "tools", *tool_task("find my tax pdf in Documents", _t4)),
        ("T5", "tools", *tool_task("put a meeting with Omar next Monday at 10 for 45 minutes on my calendar", _t5)),
        ("T6", "tools", *tool_task("what's the capital of France?", _t6)),
        ("T7", "tools", *tool_task("play something calm by Fairuz", _t7)),
        ("T8", "tools", *tool_task("write a short thank-you note in Word titled Thanks", _t8)),
        ("E1", "email", "Extract from this email. Reply with ONLY JSON with keys sender_name, sender_email, event_date (ISO yyyy-mm-dd), "
                        "event_time (24h HH:MM), location, deadline (ISO yyyy-mm-dd), action_items (list of the 3 requested actions).\n\n" + EMAIL1, chk_e1),
        ("E2", "email", "Summarize this thread as exactly 3 bullet points, each starting with '- ', covering the deadline, the budget and who owns it.\n\n" + THREAD, chk_e2),
        ("E3", "email", "My boss wrote: \"Can we meet Friday at 3pm to review the Q3 numbers?\" I can't do Friday. Write my reply: polite, "
                        "apologise, propose Tuesday at 2 pm instead, at most 80 words, no exclamation marks. Only the email body.", chk_e3),
        ("E4", "email", "Classify each email as urgent, normal or spam. Reply with ONLY JSON like {\"1\":\"urgent\",...} for all 6.\n\n" +
                        "\n".join(f"{i + 1}. {s}" for i, s in enumerate(E4)), chk_e4),
        ("E5", "email", "Rewrite formally, without contractions or slang, keeping every fact: \"hey, just a heads up that invoice 8841 is still unpaid, "
                        "pls send it over by Oct 20 thx\". Only the rewritten message.", chk_e5),
        ("O1", "office", "Write a complete Python script using python-docx that creates report.docx in the current folder with: a Heading level 1 "
                         "'Q3 Summary'; a paragraph 'Revenue grew 12% quarter over quarter.'; and a table with 3 rows and 2 columns: header row "
                         "'Region','Revenue', then 'North','1.2M', then 'South','0.9M'. Only the code.", chk_o1),
        ("O2", "office", "Excel: write one formula that sums B2:B20 only where the text in A2:A20 is North. Only the formula.", chk_o2),
        ("O3", "office", "Excel: one formula for the average of C2:C50 (blank cells are ignored by Excel's average), rounded to 1 decimal. Only the formula.", chk_o3),
        ("O4", "office", "Create a 4-slide outline about a remote work policy. Reply with ONLY JSON {\"slides\":[{\"title\":...,\"bullets\":[3 strings]}]} with distinct titles.", chk_o4),
        ("O5", "office", "Turn these notes into meeting minutes with three markdown headings: Attendees, Decisions, Action items.\n\n"
                         "Notes: Mon sync. Present: Ana, Ben, Chloe. Decided to move launch to Nov 4. Ben to fix login bug by Friday. Chloe books the venue. Ana will email the client.", chk_o5),
        ("S1", "scrape", "Convert this HTML table to a JSON array of objects with keys Product (string), Qty (integer), Price (number). Only JSON.\n\n" + HTML1, chk_s1),
        ("S2", "scrape", "Extract the products as a JSON array of {name, price (number, no $), in_stock (boolean)}. Only JSON.\n\n" + HTML2, chk_s2),
        ("S3", "scrape", "Write a Python function extract(html) using BeautifulSoup (from bs4 import BeautifulSoup) that returns a list of (text, href) tuples "
                         "for every <a> inside the <nav> element, in document order, with the text stripped of whitespace. Only the code.", chk_s3),
        ("S4", "scrape", "Extract every valid email address and phone number from the text as JSON {\"emails\":[...],\"phones\":[...]}. Only JSON.\n\n" + MESSY, chk_s4),
        ("W1", "windows", "From this tasklist output reply with ONLY JSON {\"top_process\": name of the process using the most memory, \"chrome_total_mb\": total memory of all "
                          "chrome.exe processes in megabytes, one decimal}.\n\n" + TASKLIST, chk_w1),
        ("W2", "windows", "From this systeminfo output reply with ONLY JSON {\"os\": string, \"ram_gb\": total RAM in GB one decimal, \"free_ram_gb\": available RAM in GB one "
                          "decimal, \"uptime_days\": whole days since boot, given that now is 10/8/2026, 8:15 AM}.\n\n" + SYSINFO, chk_w2),
        ("W3", "windows", "From this file listing, return a JSON array of file names larger than 5 MB (5*1024*1024 bytes) and modified after 2026-09-01, largest first. Only JSON.\n\n" + FILES, chk_w3),
        ("V1", "screen", "This is the accessibility tree of the active window. Which element id should be clicked to save this document under a new name? Reply with ONLY JSON {\"id\": number}.\n\n" + UITREE, chk_v1),
        ("V2", "screen", "Accessibility tree of an email compose window:\n" + COMPOSE + "\nGive the steps to send an email to bob@x.com as ONLY JSON {\"steps\":[{\"id\":number,\"action\":\"click\" or \"type\",\"text\":optional string}]}.", chk_v2),
        ("V3", "screen", "OCR text of an invoice (may contain OCR mistakes, fix obvious ones). Reply with ONLY JSON {\"invoice_number\": string, \"total\": number, \"due_date\": ISO yyyy-mm-dd}.\n\n" + OCR, chk_v3),
        ("R1", "reason", "Schedule 4 meetings on one day between 09:00 and 13:00 with at least 15 minutes between consecutive meetings: A is 60 min, B 30 min, C 45 min, "
                         "D 30 min. B must be before A. C cannot be first. D must be last. Reply with ONLY a JSON list of {\"name\",\"start\",\"end\"} using HH:MM.", chk_r1),
        ("R2", "reason", "Write exactly 5 lines. Line n must be 'n. ' followed by one lowercase English word that starts with the letter b and has exactly n+2 letters.", chk_r2),
        ("R3", "reason", "A train leaves at 08:40 and travels 150 km at 90 km/h, stops for 12 minutes, then travels 60 km at 72 km/h. At what time does it arrive? Reply as HH:MM.", chk_r3),
    ]
    T.append(("L1", "long", "Read this long document, then reply with ONLY JSON {\"budget_code\": ..., \"vendor_deadline\": ..., \"kickoff_room\": ...} "
                            "using the notes buried in it.\n\n" + _long_doc(), chk_l1))
    V = [("X1", "vision", "What total amount is written in the image? Reply with only the number as written.", chk_v4, _draw_total),
         ("X2", "vision", "What color is the rightmost circle in the image? Reply with one word.", chk_v5, _draw_circles)]
    return T, V


# ───────────────────────────── running ─────────────────────────────
def call(backend: str, model: str, prompt: str, images: list[str] | None = None) -> dict:
    msg = {"role": "user", "content": prompt}
    if images:
        msg["images"] = images
    body = {"model": model, "stream": False, "messages": [{"role": "system", "content": SYSTEM}, msg],
            "options": {"temperature": 0, "num_predict": 4000 if backend == "cloud" else 1500}}
    if backend == "local":
        body["think"] = False
    url = (CLOUD if backend == "cloud" else LOCAL) + "/api/chat"
    headers = {"Authorization": f"Bearer {KEY}"} if backend == "cloud" else {}
    for attempt in range(4):
        started = time.time()
        try:
            r = requests.post(url, json=body, headers=headers, timeout=240)
        except requests.RequestException as exc:
            return {"error": type(exc).__name__, "secs": round(time.time() - started, 1)}
        if r.status_code == 429 or r.status_code >= 500:
            time.sleep(15 * (attempt + 1))
            continue
        try:
            data = r.json()
        except ValueError:
            data = {}
        if r.status_code != 200:
            return {"error": f"HTTP {r.status_code}: {str(data.get('error', ''))[:100]}", "secs": round(time.time() - started, 1)}
        msg_out = data.get("message") or {}
        text = str(msg_out.get("content", ""))
        calls = msg_out.get("tool_calls") or []
        if not text.strip() and calls:      # the model answered through the native tool-call field: same answer, other channel
            fn = calls[0].get("function") or {}
            text = json.dumps({"tool": fn.get("name"), "args": fn.get("arguments") or {}})
        return {"text": text, "secs": round(time.time() - started, 1),
                "in": data.get("prompt_eval_count", 0), "out": data.get("eval_count", 0)}
    return {"error": "rate limited after retries", "secs": 0}


def done_pairs() -> set:
    if not OUT.exists():
        return set()
    return {(r["model"], r["task"]) for r in (json.loads(l) for l in OUT.read_text(encoding="utf-8").splitlines() if l.strip())}


def run(backend: str, models: list[str]) -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    tasks, vis = build_tasks()
    seen = done_pairs()
    for model in models:
        for tid, cat, prompt, check, *extra in tasks + [(*v,) for v in vis]:
            if (model, tid) in seen:
                continue
            images = None
            if cat == "vision":
                if "gemma4" not in model:
                    continue                                    # only the vision-capable model is asked
                images = [png_b64(extra[0])]
            res = call(backend, model, prompt, images)
            if "error" in res:
                row = {"model": model, "task": tid, "cat": cat, "score": 0.0, "note": res["error"], "secs": res["secs"], "in": 0, "out": 0, "raw": ""}
            else:
                try:
                    score, note = check(res["text"])
                except Exception as exc:  # noqa: BLE001
                    score, note = 0.0, f"checker error {type(exc).__name__}"
                row = {"model": model, "task": tid, "cat": cat, "score": round(score, 3), "note": note, "secs": res["secs"],
                       "in": res["in"], "out": res["out"], "raw": res["text"][:1500]}
            with OUT.open("a", encoding="utf-8") as f:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
            print(f"{model:20s} {tid:3s} {row['score']:.2f}  {row['secs']:5.1f}s  {row['note'][:60]}", flush=True)


def report() -> None:
    rows = [json.loads(l) for l in OUT.read_text(encoding="utf-8").splitlines() if l.strip()]
    models = sorted({r["model"] for r in rows})
    cats = ["tools", "email", "office", "scrape", "windows", "screen", "reason", "long"]
    print(f"{'model':20s} " + " ".join(f"{c:>7s}" for c in cats) + "   mean  strict  avg s  tokens in/out   list-cost")
    for m in models:
        rs = [r for r in rows if r["model"] == m and r["cat"] != "vision"]
        by = defaultdict(list)
        for r in rs:
            by[r["cat"]].append(r["score"])
        pi, po = PRICE.get(m, (0, 0))
        ti, to = sum(r["in"] for r in rs), sum(r["out"] for r in rs)
        print(f"{m:20s} " + " ".join(f"{sum(by[c]) / max(1, len(by[c])):7.2f}" for c in cats) +
              f"   {sum(r['score'] for r in rs) / max(1, len(rs)):.2f}   {sum(r['score'] == 1 for r in rs):2d}/{len(rs)}   {sum(r['secs'] for r in rs) / max(1, len(rs)):5.1f}"
              f"   {ti}/{to}   ${(ti * pi + to * po) / 1e6:.4f}")
    print("\nvision (gemma4 only):", {r["task"]: r["score"] for r in rows if r["cat"] == "vision"})


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--backend", choices=["cloud", "local"])
    ap.add_argument("--models", nargs="*", default=[])
    ap.add_argument("--report", action="store_true")
    a = ap.parse_args()
    if a.report:
        report()
    else:
        if a.backend == "cloud" and not KEY:
            sys.exit("OLLAMA_CLOUD_API_KEY is not set in .env")
        run(a.backend, a.models)
